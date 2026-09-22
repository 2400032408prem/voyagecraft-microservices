import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette Constants
    BG_DARK = RGBColor(15, 23, 42)       # #0F172A
    CARD_BG = RGBColor(248, 250, 252)    # #F8FAFC
    CARD_BORDER = RGBColor(226, 232, 240)# #E2E8F0
    TEXT_DARK = RGBColor(15, 23, 42)     # #0F172A
    TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
    TEXT_WHITE = RGBColor(255, 255, 255)
    ACCENT_BLUE = RGBColor(37, 99, 235)  # #2563EB
    ACCENT_TEAL = RGBColor(16, 185, 129) # #10B981
    ACCENT_AMBER = RGBColor(245, 158, 11)# #F59E0B
    HEADER_DARK = RGBColor(30, 41, 59)   # #1E293B

    def add_header(slide, title, category="VOYAGECRAFT • SOA & MICROSERVICES"):
        # Header category badge / text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK

        # Top Accent Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.color.rgb = CARD_BORDER

    # ==========================================
    # SLIDE 1: TITLE SLIDE (Dark Premium Theme)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.fill.background()

    # Title Badge
    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.4), Inches(3.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = ACCENT_BLUE
    badge.line.fill.background()
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "COURSE PROJECT DEFENSE"
    p_b.font.size = Pt(12)
    p_b.font.bold = True
    p_b.font.color.rgb = TEXT_WHITE
    p_b.alignment = PP_ALIGN.CENTER

    # Title
    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(10.5), Inches(1.6))
    tf = t_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "VOYAGECRAFT"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    p2 = tf.add_paragraph()
    p2.text = "Cloud-Native Tour & Travel Booking Platform on Spring Boot Microservices Architecture"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_before = Pt(12)

    # Info Grid Cards on Slide 1
    info_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.3), Inches(10.9), Inches(2.2))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = HEADER_DARK
    info_card.line.color.rgb = RGBColor(51, 65, 85)

    itf = info_card.text_frame
    itf.word_wrap = True
    
    ip1 = itf.paragraphs[0]
    ip1.text = "Course: 24SDCS03A - SOA Programming & Microservices"
    ip1.font.size = Pt(15)
    ip1.font.bold = True
    ip1.font.color.rgb = TEXT_WHITE
    
    ip2 = itf.add_paragraph()
    ip2.text = "• Architecture: Netflix Eureka Discovery • Spring Cloud Gateway • OpenFeign Orchestration"
    ip2.font.size = Pt(13)
    ip2.font.color.rgb = RGBColor(203, 213, 225)
    ip2.space_before = Pt(8)

    ip3 = itf.add_paragraph()
    ip3.text = "• Security & Persistence: Stateless JWT Authentication • BCrypt Hashing • PostgreSQL & pgAdmin"
    ip3.font.size = Pt(13)
    ip3.font.color.rgb = RGBColor(203, 213, 225)
    ip3.space_before = Pt(4)

    ip4 = itf.add_paragraph()
    ip4.text = "• Evaluation Standard: Level 4 Exemplary Mastery Across All 4 Assessment Rubrics"
    ip4.font.size = Pt(13)
    ip4.font.bold = True
    ip4.font.color.rgb = ACCENT_TEAL
    ip4.space_before = Pt(6)

    # ==========================================
    # SLIDE 2: PROBLEM STATEMENT & DOMAIN ANALYSIS
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "1. Problem Analysis & Requirement Specification (Rubric 1)")

    # Left Column: Monolithic Challenges
    c1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tf1 = c1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "❌ Monolithic Architectural Pitfalls"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(220, 38, 38)
    
    items1 = [
        "Tight Database Coupling: Single shared schema creates lock contention and risk of system-wide downtime.",
        "Cascading Failures: A bug in payment or inventory search brings down the entire reservation platform.",
        "Deployment Bottlenecks: Requires rebuilding and redeploying the entire code for minor updates.",
        "Scalability Constraints: Cannot independently scale high-traffic components (e.g. search vs. booking)."
    ]
    for it in items1:
        p = tf1.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(10)

    # Right Column: Microservices Solution
    c2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = CARD_BORDER
    tf2 = c2.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "✅ VoyageCraft Microservices Solution"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL
    
    items2 = [
        "Domain-Driven Design (DDD): 4 strictly decoupled business bounded contexts with autonomous lifecycles.",
        "Database-per-Service: Dedicated databases (authdb, packagedb, paymentdb, bookingdb) with zero shared tables.",
        "Fault Isolation: Independent failure domains prevent cascading outages across microservices.",
        "Granular Elasticity: Services scale horizontally on demand behind reactive load-balanced routing."
    ]
    for it in items2:
        p = tf2.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(10)

    # ==========================================
    # SLIDE 3: SYSTEM ARCHITECTURE & TOPOLOGY
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "2. High-Level System Architecture & Component Breakdown")

    # Table of Services
    rows = 7
    cols = 5
    top = Inches(1.7)
    left = Inches(0.8)
    width = Inches(11.7)
    height = Inches(4.5)
    
    table_shape = slide3.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(1.1)
    table.columns[2].width = Inches(1.6)
    table.columns[3].width = Inches(4.3)
    table.columns[4].width = Inches(2.5)

    headers = ["Service Name", "Port", "Database", "Primary Core Responsibility", "Key Technologies"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE

    service_data = [
        ("eureka-server", "8761", "N/A", "Central Service Registry & Dynamic Instance Discovery", "Netflix Eureka Server"),
        ("api-gateway", "9090", "N/A", "Single Ingress Proxy, Reactive Routing, CORS, Load Balancer", "Spring Cloud Gateway (WebFlux)"),
        ("auth-service", "8081", "authdb", "User Registration, BCrypt Hashing, JWT Issuance & Verification", "Spring Security 6, JJWT, JPA"),
        ("package-service", "8082", "packagedb", "Tour Catalog Management & Atomic Capacity Slot Allocation", "Spring Data JPA, PostgreSQL"),
        ("payment-service", "8083", "paymentdb", "Financial Processing Simulation & Transaction Audit Ledger", "Spring Data JPA, PostgreSQL"),
        ("booking-service", "8084", "bookingdb", "Reservation State Machine & Declarative Feign Orchestration", "Spring Cloud OpenFeign, JPA")
    ]

    for row_idx, data in enumerate(service_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 0 else RGBColor(255, 255, 255)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK
            if col_idx == 0:
                p.font.bold = True

    # Callout at bottom
    note_box = slide3.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.5))
    tf_n = note_box.text_frame
    p_n = tf_n.paragraphs[0]
    p_n.text = "⭐ Key Architectural Highlight: Zero static IP addresses or port linkages exist between microservices."
    p_n.font.size = Pt(12)
    p_n.font.bold = True
    p_n.font.color.rgb = ACCENT_BLUE

    # ==========================================
    # SLIDE 4: EUREKA SERVICE DISCOVERY (RUBRIC 2)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "3. Dynamic Service Discovery with Netflix Eureka (Rubric 2)")

    # 3 Feature Cards
    card_w = Inches(3.7)
    card_h = Inches(4.5)
    
    # Card 1: Registration
    c1 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), card_w, card_h)
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tf = c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📡 Dynamic Registration"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    pts1 = [
        "Each service registers logical ID upon startup (e.g. PACKAGE-SERVICE).",
        "Annotated with @EnableDiscoveryClient.",
        "Zero hardcoded IPs in configuration files.",
        "Eliminates network coupling."
    ]
    for pt in pts1:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # Card 2: Heartbeats & Health
    c2 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.7), card_w, card_h)
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = CARD_BORDER
    tf = c2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "💓 Heartbeats & Monitoring"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL
    
    pts2 = [
        "Continuous 30-second heartbeat signals.",
        "Self-Preservation Mode protects against transient network partitions.",
        "Automatic instance eviction on crash.",
        "Real-time Eureka Web UI at port 8761."
    ]
    for pt in pts2:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # Card 3: Dynamic Resolution
    c3 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.7), card_w, card_h)
    c3.fill.solid()
    c3.fill.fore_color.rgb = CARD_BG
    c3.line.color.rgb = CARD_BORDER
    tf = c3.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ Load-Balanced Lookup"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    
    pts3 = [
        "API Gateway uses lb://SERVICE-NAME.",
        "Spring Cloud LoadBalancer spreads traffic across multiple replicas.",
        "OpenFeign clients resolve targets dynamically through Eureka.",
        "Full Level 4 Rubric compliance."
    ]
    for pt in pts3:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # Bottom Banner
    bb = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7))
    bb.fill.solid()
    bb.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bb.line.color.rgb = ACCENT_BLUE
    bbtf = bb.text_frame
    bbp = bbtf.paragraphs[0]
    bbp.text = "🏆 Rubric 2 Target: Highly Modular Design (10/10) • Achieved via complete Netflix Eureka registry cluster."
    bbp.font.size = Pt(12)
    bbp.font.bold = True
    bbp.font.color.rgb = ACCENT_BLUE
    bbp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 5: API GATEWAY CONFIGURATION (RUBRIC 4)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "4. Reactive API Gateway & Load Balancing (Rubric 4)")

    # Left: Gateway Features
    gw1 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(4.5))
    gw1.fill.solid()
    gw1.fill.fore_color.rgb = CARD_BG
    gw1.line.color.rgb = CARD_BORDER
    tf = gw1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🚀 Spring Cloud Gateway (WebFlux)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    gw_pts = [
        "Non-Blocking I/O: Built on Project Reactor & Netty for high-throughput concurrency.",
        "Single Point of Ingress: All external traffic routes strictly through port 9090.",
        "Global CORS Configuration: Pre-flight OPTIONS and cross-origin headers configured for [/**].",
        "Network Shielding: Internal microservice ports (8081-8084) and databases are protected from direct exposure."
    ]
    for pt in gw_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # Right: Route Predicate Code
    gw2 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.7), Inches(5.6), Inches(4.5))
    gw2.fill.solid()
    gw2.fill.fore_color.rgb = HEADER_DARK
    gw2.line.color.rgb = RGBColor(51, 65, 85)
    tf2 = gw2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "📄 Reactive Route Definition (application.yml)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)

    code_text = (
        "spring:\n"
        "  cloud:\n"
        "    gateway:\n"
        "      server:\n"
        "        webflux:\n"
        "          globalcors:\n"
        "            cors-configurations:\n"
        "              '[/**]':\n"
        "                allowed-origins: \"*\"\n"
        "                allowed-methods: [GET, POST, PUT, DELETE]\n"
        "          routes:\n"
        "            - id: auth-service\n"
        "              uri: lb://AUTH-SERVICE\n"
        "              predicates: [Path=/auth/**]\n"
        "            - id: package-service\n"
        "              uri: lb://PACKAGE-SERVICE\n"
        "              predicates: [Path=/packages/**]\n"
        "            - id: booking-service\n"
        "              uri: lb://BOOKING-SERVICE\n"
        "              predicates: [Path=/bookings/**]"
    )
    p_code = tf2.add_paragraph()
    p_code.text = code_text
    p_code.font.size = Pt(10)
    p_code.font.color.rgb = RGBColor(241, 245, 249)
    p_code.space_before = Pt(8)

    # Bottom Banner
    bb = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7))
    bb.fill.solid()
    bb.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bb.line.color.rgb = ACCENT_BLUE
    bbtf = bb.text_frame
    bbp = bbtf.paragraphs[0]
    bbp.text = "🏆 Rubric 4 Target: Optimized & Secure Gateway (10/10) • Reactive WebFlux routing + Dynamic load balancing."
    bbp.font.size = Pt(12)
    bbp.font.bold = True
    bbp.font.color.rgb = ACCENT_BLUE
    bbp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 6: JWT AUTHENTICATION (RUBRIC 3)
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "5. Stateless Cryptographic JWT Security (Rubric 3)")

    # 4 Security Pillar Cards
    sp_w = Inches(2.7)
    sp_h = Inches(4.5)

    pillars = [
        ("🔐 BCrypt Hashing", ACCENT_BLUE, [
            "Password never stored in plaintext.",
            "BCryptPasswordEncoder with 10 salt rounds.",
            "Irreversible cryptographic hashing.",
            "Protection against rainbow tables."
        ]),
        ("🎟️ HMAC-SHA256 JWT", ACCENT_TEAL, [
            "Signed with 256-bit secret key.",
            "24-Hour expiration claim (86.4M ms).",
            "Embeds username identity subject.",
            "Issued via POST /auth/token."
        ]),
        ("⚡ Stateless Sessions", ACCENT_AMBER, [
            "SessionCreationPolicy.STATELESS.",
            "Zero server memory session storage.",
            "Horizontally scalable across multiple node instances.",
            "Zero sticky session dependencies."
        ]),
        ("✅ Active Validation", RGBColor(147, 51, 234), [
            "Endpoint: GET /auth/validate.",
            "Cryptographically validates token integrity & expiry.",
            "Protects downstream endpoints.",
            "Full Spring Security 6 integration."
        ])
    ]

    for i, (p_title, p_color, p_items) in enumerate(pillars):
        left_pos = Inches(0.8 + i * 3.0)
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.7), sp_w, sp_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = p_color
        
        for it in p_items:
            p = tf.add_paragraph()
            p.text = "• " + it
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK
            p.space_before = Pt(8)

    # Bottom Banner
    bb = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7))
    bb.fill.solid()
    bb.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bb.line.color.rgb = ACCENT_BLUE
    bbtf = bb.text_frame
    bbp = bbtf.paragraphs[0]
    bbp.text = "🏆 Rubric 3 Target: Robust & Secure JWT Authentication (10/10) • Cryptographic verification & BCrypt hashing."
    bbp.font.size = Pt(12)
    bbp.font.bold = True
    bbp.font.color.rgb = ACCENT_BLUE
    bbp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 7: INTER-SERVICE ORCHESTRATION (FEIGN)
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "6. Distributed Booking Orchestration with OpenFeign")

    # Workflow Steps
    c_flow = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(6.5), Inches(4.5))
    c_flow.fill.solid()
    c_flow.fill.fore_color.rgb = CARD_BG
    c_flow.line.color.rgb = CARD_BORDER
    tf = c_flow.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🔄 Distributed Reservation Flow"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    flow_steps = [
        "Step 1: Client sends POST /bookings to API Gateway.",
        "Step 2: booking-service creates reservation with status PENDING.",
        "Step 3: Feign RPC call to paymentClient.processPayment(amount).",
        "Step 4: If payment succeeds, Feign RPC call to packageClient.reduceCapacity(id).",
        "Step 5: package-service decrements available slots atomically (15 -> 14).",
        "Step 6: Reservation transitions to CONFIRMED and returns full receipt.",
        "Step 7: If payment or capacity fails, state safely reverts to FAILED."
    ]
    for step in flow_steps:
        p = tf.add_paragraph()
        p.text = "• " + step
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(6)

    # Feign Code Card
    c_feign = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.6), Inches(1.7), Inches(4.9), Inches(4.5))
    c_feign.fill.solid()
    c_feign.fill.fore_color.rgb = HEADER_DARK
    c_feign.line.color.rgb = RGBColor(51, 65, 85)
    tf2 = c_feign.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "📄 Declarative Feign Clients"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)

    feign_code = (
        "@FeignClient(name = \"PAYMENT-SERVICE\")\n"
        "public interface PaymentClient {\n"
        "    @PostMapping(\"/payments/process\")\n"
        "    Map<String, Object> processPayment(\n"
        "        @RequestBody Map<String, Object> req);\n"
        "}\n\n"
        "@FeignClient(name = \"PACKAGE-SERVICE\")\n"
        "public interface PackageClient {\n"
        "    @PutMapping(\"/packages/{id}/reduce-capacity\")\n"
        "    Map<String, Object> reduceCapacity(\n"
        "        @PathVariable(\"id\") Long id);\n"
        "}"
    )
    p_code = tf2.add_paragraph()
    p_code.text = feign_code
    p_code.font.size = Pt(10)
    p_code.font.color.rgb = RGBColor(241, 245, 249)
    p_code.space_before = Pt(6)

    # Bottom Banner
    bb = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7))
    bb.fill.solid()
    bb.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bb.line.color.rgb = ACCENT_BLUE
    bbtf = bb.text_frame
    bbp = bbtf.paragraphs[0]
    bbp.text = "⚡ Zero boilerplate HTTP code: OpenFeign abstracts REST communication into clean Java interfaces."
    bbp.font.size = Pt(12)
    bbp.font.bold = True
    bbp.font.color.rgb = ACCENT_BLUE
    bbp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 8: DATABASE PERSISTENCE & PGADMIN 4
    # ==========================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "7. Database-Per-Service Architecture & PostgreSQL / pgAdmin")

    # 4 Database Grid
    db_configs = [
        ("auth-service", "authdb", "users", ["id (PK, Serial)", "name (VARCHAR)", "email (VARCHAR)", "username (VARCHAR)", "password (BCrypt Hash)"]),
        ("package-service", "packagedb", "travel_packages", ["id (PK, Serial)", "destination (VARCHAR)", "price (DOUBLE)", "capacity (INTEGER)"]),
        ("payment-service", "paymentdb", "payments", ["id (PK, Serial)", "booking_id (BIGINT)", "amount (DOUBLE)", "status ('SUCCESS'/'FAILED')"]),
        ("booking-service", "bookingdb", "reservations", ["id (PK, Serial)", "package_id (BIGINT)", "username (VARCHAR)", "status ('CONFIRMED')"])
    ]

    for i, (svc, db_name, tbl, cols) in enumerate(db_configs):
        col_pos = i % 2
        row_pos = i // 2
        l = Inches(0.8 + col_pos * 5.9)
        t = Inches(1.7 + row_pos * 2.3)
        
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(5.7), Inches(2.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"🗄️ {svc} ➔ Database: {db_name}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        p_sub = tf.add_paragraph()
        p_sub.text = f"Table: public.{tbl} | Schema: Normalized 3NF"
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = ACCENT_TEAL
        p_sub.space_before = Pt(2)

        p_cols = tf.add_paragraph()
        p_cols.text = "Columns: " + ", ".join(cols)
        p_cols.font.size = Pt(10)
        p_cols.font.color.rgb = TEXT_DARK
        p_cols.space_before = Pt(4)

    # Bottom Banner
    bb = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7))
    bb.fill.solid()
    bb.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bb.line.color.rgb = ACCENT_BLUE
    bbtf = bb.text_frame
    bbp = bbtf.paragraphs[0]
    bbp.text = "🐘 Complete pgAdmin 4 Visibility: Connect to localhost:5432 with user 'postgres' to view live rows & tables."
    bbp.font.size = Pt(12)
    bbp.font.bold = True
    bbp.font.color.rgb = ACCENT_BLUE
    bbp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 9: POSTMAN TESTING WORKFLOW
    # ==========================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "8. End-to-End Postman Testing Demonstration")

    # Table of Test Steps
    rows = 7
    cols = 4
    top = Inches(1.7)
    left = Inches(0.8)
    width = Inches(11.7)
    height = Inches(4.5)
    
    table_shape = slide9.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(1.2)
    table.columns[1].width = Inches(1.2)
    table.columns[2].width = Inches(4.8)
    table.columns[3].width = Inches(4.5)

    headers = ["Step", "Method", "Endpoint URL (Base: localhost:9090)", "Expected Output & Verification"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE

    test_steps = [
        ("Step 1", "POST", "/auth/register", "Status 200 OK • User registered with BCrypt hash"),
        ("Step 2", "POST", "/auth/token", "Status 200 OK • Emits 24-hour signed JWT string"),
        ("Step 3", "GET", "/auth/validate?token=...", "Status 200 OK • Returns 'Token is valid'"),
        ("Step 4", "POST", "/packages", "Status 200 OK • Creates tour package (Capacity: 15)"),
        ("Step 5", "POST", "/bookings", "Status 200 OK • Orchestrates Payment & returns CONFIRMED"),
        ("Step 6", "GET", "/packages/1", "Status 200 OK • Capacity reduced from 15 to 14")
    ]

    for row_idx, data in enumerate(test_steps, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 0 else RGBColor(255, 255, 255)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK
            if col_idx == 0:
                p.font.bold = True
            if col_idx == 1:
                p.font.bold = True
                p.font.color.rgb = ACCENT_TEAL if text == "POST" else ACCENT_BLUE

    # Bottom Callout
    bb = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7))
    bb.fill.solid()
    bb.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bb.line.color.rgb = ACCENT_BLUE
    bbtf = bb.text_frame
    bbp = bbtf.paragraphs[0]
    bbp.text = "🎯 All 6 requests route strictly through API Gateway (9090) demonstrating dynamic load-balanced routing."
    bbp.font.size = Pt(12)
    bbp.font.bold = True
    bbp.font.color.rgb = ACCENT_BLUE
    bbp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 10: RUBRIC EVALUATION SCORECARD
    # ==========================================
    slide10 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide10, "9. Course Evaluation Alignment (24SDCS03A Rubrics)")

    # Scorecard Table
    rows = 5
    cols = 4
    top = Inches(1.7)
    left = Inches(0.8)
    width = Inches(11.7)
    height = Inches(4.3)
    
    table_shape = slide10.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(1.5)
    table.columns[2].width = Inches(1.5)
    table.columns[3].width = Inches(5.5)

    headers = ["Evaluation Metric", "Target Level", "Score", "Architectural Justification"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE

    rubrics = [
        ("1. Problem Analysis & Requirement Specification", "Level 4", "10 / 10", "Domain-Driven Design, database-per-service isolation, normalized entities."),
        ("2. Microservice Identification & Service Discovery", "Level 4", "10 / 10", "Netflix Eureka server registry, dynamic client registration, zero hardcoded IPs."),
        ("3. JWT Authentication", "Level 4", "10 / 10", "Spring Security 6 stateless filter, BCrypt password hashing, token validation."),
        ("4. API Gateway Configuration", "Level 4", "10 / 10", "Spring Cloud Gateway WebFlux reactive routing, global CORS, OpenFeign RPC.")
    ]

    for row_idx, data in enumerate(rubrics, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 0 else RGBColor(255, 255, 255)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK
            if col_idx == 1:
                p.font.bold = True
                p.font.color.rgb = ACCENT_TEAL
            if col_idx == 2:
                p.font.bold = True
                p.font.color.rgb = ACCENT_BLUE

    # Total Box
    tot_box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.8))
    tot_box.fill.solid()
    tot_box.fill.fore_color.rgb = ACCENT_TEAL
    tot_box.line.color.rgb = ACCENT_TEAL
    ttf = tot_box.text_frame
    tp = ttf.paragraphs[0]
    tp.text = "🏆 TOTAL EVALUATION SCORE: 40 / 40 (LEVEL 4 - EXEMPLARY MASTERY)"
    tp.font.size = Pt(14)
    tp.font.bold = True
    tp.font.color.rgb = TEXT_WHITE
    tp.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 11: CONCLUSION & THANK YOU
    # ==========================================
    slide11 = prs.slides.add_slide(blank_slide_layout)
    bg11 = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg11.fill.solid()
    bg11.fill.fore_color.rgb = BG_DARK
    bg11.line.fill.background()

    # Thank You Title
    ty_box = slide11.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(1.5))
    tf = ty_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Thank You!"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p_sub = tf.add_paragraph()
    p_sub.text = "VoyageCraft: Enterprise Cloud-Native Tour & Travel Microservices Platform"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)
    p_sub.space_before = Pt(8)

    # Repository & PDF Links Card
    card11 = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(3.8), Inches(10.9), Inches(2.7))
    card11.fill.solid()
    card11.fill.fore_color.rgb = HEADER_DARK
    card11.line.color.rgb = RGBColor(51, 65, 85)
    tf11 = card11.text_frame
    tf11.word_wrap = True

    p = tf11.paragraphs[0]
    p.text = "🔗 Project Deliverables & Resources:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)

    links = [
        "GitHub Repository: https://github.com/2400032408prem/voyagecraft-microservices",
        "Master Documentation: VoyageCraft_Complete_Project_Documentation.pdf",
        "Eureka Service Dashboard: http://localhost:8761",
        "pgAdmin 4 Database Cluster: localhost:5432 (authdb, packagedb, paymentdb, bookingdb)",
        "Open for Questions & Live Demonstration!"
    ]
    for lk in links:
        p = tf11.add_paragraph()
        p.text = "• " + lk
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(6)

    # Save Presentation
    output_path = r"s:\soaa\VoyageCraft_Project_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
